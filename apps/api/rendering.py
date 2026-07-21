"""Tailored resume rendering: clone the master .docx, swap in the plan's
SELECTION, convert to PDF with Word COM, verify the rendered PDF.

Ported from the InternshipApplier (its clone-of-master renderer): formatting
is identical to the master by construction; descriptions render hyphen-free
while dates/headers keep theirs; the resume is never "ready" from the XML
alone — verify_render() checks the produced PDF.

Windows + Word only. The deployed Linux API never imports win32com; callers
gate on render_available(). Long-term this moves to the runner (spec §7.1).
"""
import re
from pathlib import Path

PLACEHOLDER_MARKERS = ("<add", "TODO", "{{", "}}", "lorem")


class RenderError(Exception):
    pass


def render_available() -> bool:
    try:
        import win32com.client  # noqa: F401
        from docx import Document  # noqa: F401
        return True
    except Exception:
        return False


def _dehyph(s: str) -> str:
    t = s.replace("-", " ")
    while "  " in t:
        t = t.replace("  ", " ")
    return t


def _qn(tag):
    from docx.oxml.ns import qn
    return qn(tag)


def _copy_rpr(src_run_el):
    import copy as _c
    rpr = src_run_el.find(_qn("w:rPr")) if src_run_el is not None else None
    return _c.deepcopy(rpr) if rpr is not None else None


def _set_para_runs(p, text, metric, template_rpr, bold_rpr):
    import copy as _c
    for r in list(p.runs):
        r._r.getparent().remove(r._r)

    def add(t, bold):
        if not t:
            return
        r = p.add_run(t)
        rpr = bold_rpr if bold and bold_rpr is not None else template_rpr
        if rpr is not None:
            r._r.insert(0, _c.deepcopy(rpr))
        if bold:
            r.bold = True
    text = _dehyph(" ".join(text.split()))
    metric = _dehyph(metric or "")
    if metric and metric in text:
        pre, _, post = text.partition(metric)
        add(pre, False)
        add(metric, True)
        add(post, False)
    else:
        add(text, False)


def _doc_structure(doc):
    paras = doc.paragraphs
    idx = {p.text.strip(): i for i, p in enumerate(paras)
           if p.text.strip() in ("Experience", "Projects", "Accomplishments")}
    if not {"Experience", "Projects", "Accomplishments"} <= set(idx):
        raise RenderError("master docx section headings not found")

    def is_bullet(p):
        return p.style is not None and "List" in (p.style.name or "")

    roles, i = [], idx["Experience"] + 1
    while i < idx["Projects"]:
        if not paras[i].text.strip():
            i += 1
            continue
        org_p, title_p = paras[i], paras[i + 1]
        j = i + 2
        bullets = []
        while j < idx["Projects"] and is_bullet(paras[j]):
            bullets.append(paras[j])
            j += 1
        roles.append({"org": org_p, "title": title_p, "bullets": bullets})
        i = j
    projects, i = [], idx["Projects"] + 1
    while i < idx["Accomplishments"]:
        if not paras[i].text.strip():
            i += 1
            continue
        name_p = paras[i]
        j = i + 1
        bullets = []
        while j < idx["Accomplishments"] and is_bullet(paras[j]):
            bullets.append(paras[j])
            j += 1
        projects.append({"name": name_p, "bullets": bullets})
        i = j
    return roles, projects


def _replace_block_bullets(block_bullets, selected, pool_bullets):
    import copy as _c

    from docx.text.paragraph import Paragraph
    template_el = block_bullets[0]._p
    template_rpr = _copy_rpr(template_el.find(_qn("w:r")))
    bold_rpr = None
    for bp in block_bullets:
        for r in bp.runs:
            if r.bold:
                bold_rpr = _copy_rpr(r._r)
                break
        if bold_rpr is not None:
            break
    parent = template_el.getparent()
    anchor_idx = parent.index(template_el)
    for bp in block_bullets:
        parent.remove(bp._p)
    for offset, b in enumerate(selected):
        el = _c.deepcopy(template_el)
        for r in el.findall(_qn("w:r")):
            el.remove(r)
        parent.insert(anchor_idx + offset, el)
        p = Paragraph(el, block_bullets[0]._parent)
        _set_para_runs(p, pool_bullets[b]["text"], pool_bullets[b].get("metric", ""),
                       template_rpr, bold_rpr if bold_rpr is not None else template_rpr)


def retailor_pool_docx(master_docx: Path, pool: dict, plan: dict, out_path: Path):
    import copy as _c

    from docx import Document
    doc = Document(str(master_docx))
    roles_s, projects_s = _doc_structure(doc)
    if len(roles_s) != len(pool["experience"]):
        raise RenderError(f"master docx has {len(roles_s)} roles; pool has "
                          f"{len(pool['experience'])}")

    sp = doc.paragraphs[2]
    template_rpr = _copy_rpr(sp.runs[0]._r if sp.runs else None)
    for r in list(sp.runs):
        r._r.getparent().remove(r._r)
    summary = _dehyph(" ".join(plan["summary_text"].split()))
    bold_part = _dehyph(plan.get("summary_bold") or "")

    def add_sum(t, bold=False):
        if not t:
            return
        r = sp.add_run(t)
        if template_rpr is not None:
            r._r.insert(0, _c.deepcopy(template_rpr))
        r.bold = bold
    if bold_part and bold_part in summary:
        pre, _, post = summary.partition(bold_part)
        add_sum(pre)
        add_sum(bold_part, True)
        add_sum(post)
    else:
        add_sum(summary)

    for want in plan["skills"]:
        target = None
        for p in doc.paragraphs:
            if p.runs and p.runs[0].text.strip().rstrip(":") == want["label"]:
                target = p
                break
        if target is None:
            raise RenderError(f"skills line not found: {want['label']}")
        target.runs[-1].text = ", ".join(want["items"])

    for block, role_pool, entry in zip(roles_s, pool["experience"], plan["experience"]):
        org_p = block["org"]
        if org_p.runs:
            org_p.runs[0].text = role_pool["org"]
            for extra in org_p.runs[1:]:
                extra.text = ""
        title_p = block["title"]
        if title_p.runs:
            title_p.runs[0].text = role_pool["title"]
            if len(title_p.runs) > 1:
                title_p.runs[-1].text = "\t" + role_pool.get("dates", "")
        by_id = {b["id"]: b for b in role_pool["bullets"]}
        _replace_block_bullets(block["bullets"], entry["bullets"], by_id)

    pool_projects = {pj["id"]: pj for pj in pool["projects"]}
    if len(plan["projects"]) != len(projects_s):
        raise RenderError(f"plan selects {len(plan['projects'])} projects; the "
                          f"master docx has {len(projects_s)} slots")
    for block, entry in zip(projects_s, plan["projects"]):
        pj = pool_projects[entry["id"]]
        name_p = block["name"]
        if name_p.runs:
            name_p.runs[0].text = pj["name"]
            name_p.runs[-1].text = "  |  " + pj.get("stack", "")
        by_id = {b["id"]: b for b in pj["bullets"]}
        _replace_block_bullets(block["bullets"], entry["bullets"], by_id)

    linkedin = pool["contact"].get("linkedin", "")
    if linkedin:
        fixed_url = "https://www." + linkedin.removeprefix("www.")
        for rel in doc.part.rels.values():
            if rel.is_external and rel.reltype.endswith("/hyperlink") \
                    and rel.target_ref.rstrip("/").endswith("linkedin.com/in"):
                rel._target = fixed_url

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def docx_to_pdf(docx_path: Path, pdf_path: Path):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
        try:
            doc.SaveAs2(str(pdf_path.resolve()), FileFormat=17)  # wdFormatPDF
        finally:
            doc.Close(False)
    except Exception as e:
        raise RenderError(f"Word COM PDF conversion failed: {e}") from e
    finally:
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


def pdf_text(pdf_path: Path):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        chunks = [page.get_textpage().get_text_bounded() for page in pdf]
        return len(pdf), "\n".join(chunks)
    finally:
        pdf.close()


def _norm(s: str) -> str:
    s = s.replace("’", "'").replace("‘", "'").replace("—", "-") \
         .replace("–", "-").replace("·", " ")
    return re.sub(r"[^a-z0-9%~+]+", " ", s.lower()).strip()


def verify_render(pdf_path: Path, expected_pages: int, required_texts: list,
                  contact_bits: list) -> list:
    failures = []
    pages, text = pdf_text(pdf_path)
    if pages != expected_pages:
        failures.append(f"page count {pages} != expected {expected_pages}")
    hay = _norm(text)
    for bit in contact_bits:
        if _norm(bit) not in hay:
            failures.append(f"contact element missing: {bit!r}")
    for req in required_texts:
        probe = _norm(req)[:70]
        if probe and probe not in hay:
            failures.append(f"content missing: {req[:60]!r}...")
    low = text.lower()
    for marker in PLACEHOLDER_MARKERS:
        if marker.lower() in low:
            failures.append(f"placeholder marker leaked: {marker!r}")
    return failures


def render_job_pdf(master_docx: Path, pool: dict, plan: dict,
                   work_dir: Path) -> tuple[bytes, list]:
    """Full pipeline: clone+tailor -> PDF -> verify. Returns (pdf_bytes,
    failures). failures empty = the render passed every check."""
    work_dir.mkdir(parents=True, exist_ok=True)
    docx_path = work_dir / "tailored.docx"
    pdf_path = work_dir / "tailored.pdf"
    retailor_pool_docx(master_docx, pool, plan, docx_path)
    docx_to_pdf(docx_path, pdf_path)

    required = [plan["summary_text"]]
    roles_by_id = {r["id"]: r for r in pool["experience"]}
    for entry in plan["experience"]:
        role = roles_by_id[entry["id"]]
        required += [role["org"], role["title"]]
        by_id = {b["id"]: b for b in role["bullets"]}
        required += [by_id[bid]["text"] for bid in entry["bullets"]]
    pj_by_id = {pj["id"]: pj for pj in pool["projects"]}
    for entry in plan["projects"]:
        pj = pj_by_id[entry["id"]]
        by_id = {b["id"]: b for b in pj["bullets"]}
        required += [pj["name"]] + [by_id[bid]["text"] for bid in entry["bullets"]]
    c = pool["contact"]
    failures = verify_render(pdf_path, 2, required,
                             [c.get("email", ""), c.get("phone", "")])
    return pdf_path.read_bytes(), failures
